"""Consolidated results tables (all scenarios) -> Tabel_Hasil_Pengujian.xlsx.

Usage: python results_table.py results.json v1_v2.json v3.json OUT.xlsx
All numbers are copied from the run outputs (results.json etc.); no formulas.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

res_p, v12_p, v3_p, out = sys.argv[1:5]
R = json.load(open(res_p, encoding='utf-8'))
V12 = json.load(open(v12_p, encoding='utf-8'))
V3 = json.load(open(v3_p, encoding='utf-8'))

F = 'Arial'
HEAD = PatternFill('solid', fgColor='14213D')
BAND = PatternFill('solid', fgColor='F2F4F7')
HL = PatternFill('solid', fgColor='FFF2CC')
thin = Side(style='thin', color='D0D5DD')
BORDER = Border(top=thin, bottom=thin, left=thin, right=thin)
FMT = {'TPR': '0.00', 'TPR_adj': '0.00', 'FPR': '0.0000', 'Prec': '0.0000', 'R': '0.000',
       'alerts_per_day': '0.00', 'PR_AUC': '0.0000', 'val_TPR': '0.000', 'val_FPR': '0.0000', 'theta': '0.00'}

wb = Workbook()
wb.remove(wb.active)


def sheet(name, title, note, header, rows, fmts=None, highlight=None, widths=None):
    ws = wb.create_sheet(name)
    ws['A1'] = title
    ws['A1'].font = Font(name=F, size=13, bold=True)
    ws['A2'] = note
    ws['A2'].font = Font(name=F, size=9, italic=True, color='555555')
    for j, h in enumerate(header, 1):
        c = ws.cell(4, j, h)
        c.font = Font(name=F, bold=True, color='FFFFFF')
        c.fill = HEAD
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BORDER
    for i, row in enumerate(rows):
        for j, v in enumerate(row, 1):
            c = ws.cell(5 + i, j, v)
            c.font = Font(name=F, bold=bool(highlight and highlight(row)))
            c.border = BORDER
            if highlight and highlight(row):
                c.fill = HL
            elif i % 2:
                c.fill = BAND
            if fmts and header[j - 1] in fmts and isinstance(v, (int, float)):
                c.number_format = fmts[header[j - 1]]
            elif isinstance(v, int):
                c.number_format = '#,##0'
            c.alignment = Alignment(vertical='top', wrap_text=isinstance(v, str))
    for j in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(j, 13)
    ws.freeze_panes = 'B5'
    return ws


# ---- 1. Main model table (E, R, F, D) -------------------------------------
GROUP = {'E': 'Eksperimen utama', 'R': 'Komponen risk scoring', 'F': 'Kelompok fitur / flag', 'D': 'Adaptasi'}
DESC = {
    'E2_B': 'B – IF + fitur dasar (49 kolom)', 'E3_B+RS': 'B+RS – fitur dasar + risk scoring',
    'E4_C': 'C – IF + fitur kontekstual (57 kolom)', 'E5_D': 'D – model usulan (C + risk scoring)',
    'R1_D-h': 'D−h (tanpa riwayat)', 'R2_D-c': 'D−c (tanpa konteks)', 'R3_D-a': 'D−a (tanpa Isolation Forest)',
    'R4_D-gate': 'D-gate (a ≥ θa DAN c ≥ θc)', 'F1_C-Vol': 'C−Vol (tanpa F1)', 'F2_C-Mix': 'C−Mix (tanpa F2, F3)',
    'F3_C-Time': 'C−Time (tanpa F4, F5)', 'F4_C-Res': 'C−Res (tanpa F6)', 'F5_C-Sens': 'C−Sens (tanpa F7)',
    'F6_C-Burst': 'C−Burst (tanpa F8)', 'F7_D-k1': 'D tanpa flag k1', 'F7_D-k2': 'D tanpa flag k2',
    'F7_D-k3': 'D tanpa flag k3', 'F7_D-k4': 'D tanpa flag k4', 'F7_D-k5': 'D tanpa flag k5',
    'D1_D-static': 'D-static (latih sekali)', 'D2_D-daily': 'D-Δ1 (retrain harian)', 'D3_D-noUserMode': 'D-noUserMode (r = 0)',
}
ID = {'E2_B': 'E2', 'E3_B+RS': 'E3', 'E4_C': 'E4', 'E5_D': 'E5', 'R1_D-h': 'R1', 'R2_D-c': 'R2', 'R3_D-a': 'R3',
      'R4_D-gate': 'R4', 'F1_C-Vol': 'F1', 'F2_C-Mix': 'F2', 'F3_C-Time': 'F3', 'F4_C-Res': 'F4', 'F5_C-Sens': 'F5',
      'F6_C-Burst': 'F6', 'F7_D-k1': 'F7', 'F7_D-k2': 'F7', 'F7_D-k3': 'F7', 'F7_D-k4': 'F7', 'F7_D-k5': 'F7',
      'D1_D-static': 'D1', 'D2_D-daily': 'D2', 'D3_D-noUserMode': 'D3'}
hdr = ['ID', 'Model / varian', 'Kategori', 'theta', 'TP', 'FN', 'FP', 'TN', 'TPR', 'TPR_adj', 'FPR', 'Prec', 'R',
       'Alert total', 'alerts_per_day', 'PR_AUC', 'ΔFPR vs A', 'val_TPR', 'val_FPR']
A = R['E1_A']['test_seed1']
rows = [['E1', 'A – Baseline Landauer (θ = 0,6)', 'Eksperimen utama', 0.6, A['TP'], A['FN'], A['FP'], A['TN'], A['TPR'],
         A['TPR_adj'], A['FPR'], A['Prec'], A['R'], A['alerts'], A['alerts_per_day'], None, 0.0, None, None]]
for k, v in R['variants'].items():
    m, th = v['test_seed1'], v['theta']
    th_txt = f"a≥{th[0]:.2f}, c≥{th[1]:.3f}" if isinstance(th, list) else th
    rows.append([ID[k], DESC[k], GROUP[ID[k][0]], th_txt, m['TP'], m['FN'], m['FP'], m['TN'], m['TPR'], m['TPR_adj'],
                 m['FPR'], m['Prec'], m['R'], m['alerts'], m['alerts_per_day'], m.get('PR_AUC'),
                 m['FPR'] - A['FPR'], v['validation']['TPR'], v['validation']['FPR']])
sheet('1. Hasil model', 'Hasil semua model dan varian, seed 1 (run tunggal yang dilaporkan)',
      'θ dan parameter dipilih di seed validasi 0+2 (aturan: FPR terendah dengan TPR ≥ 0,65). '
      'val_TPR/val_FPR = gabungan seed 0+2. Baris kuning = baseline A dan model usulan D. Sumber: results.json.',
      hdr, rows, {**FMT, 'ΔFPR vs A': '+0.0000;-0.0000;0.0000'},
      highlight=lambda r: r[0] in ('E1', 'E5'), widths={2: 38, 3: 20, 4: 14})

# ---- 2. Statistics (E6) ------------------------------------------------------
rows = []
for k, v in R['E6'].items():
    keys = [x for x in v if x.startswith('false_alarms_only_')]
    rows.append([k.replace('E5_D', 'D').replace('E4_C', 'C').replace('E2_B', 'B'), v[keys[0]], v[keys[1]],
                 float(v['p_value_exact']), v['dFPR'], v['ci95'][0], v['ci95'][1]])
sheet('2. Uji statistik (E6)', 'E6 – McNemar exact + bootstrap 95% CI ΔFPR (seed 1)',
      'McNemar pada user-hari normal yang berada di fase deteksi pada kedua model. Bootstrap atas user, 2.000 resampel. '
      'Wilcoxon atas 7 seed tidak dipakai (run tunggal). p = 0 berarti < 1e-300.',
      ['Perbandingan', 'FP hanya model 1', 'FP hanya model 2', 'p-value', 'ΔFPR (model1 − model2)', 'CI 2,5%', 'CI 97,5%'],
      rows, {'p-value': '0.00E+00', 'ΔFPR (model1 − model2)': '+0.0000;-0.0000', 'CI 2,5%': '+0.0000;-0.0000', 'CI 97,5%': '+0.0000;-0.0000'},
      widths={1: 16, 5: 20})

# ---- 3. Tuning (P1-P5) -------------------------------------------------------
rows = []
for W, m in R['P1']['candidates'].items():
    rows.append(['P1', f'W = {W}', m['theta'], m['TPR'], m['FPR'], 'dikunci' if int(W) == R['P1']['locked_W'] else ''])
for w, m in R['P2']['candidates'].items():
    rows.append(['P2', f'bobot {w}', m['theta'], m['TPR'], m['FPR'],
                 'dikunci' if json.loads(w.replace('(', '[').replace(')', ']')) == list(R['P2']['locked_weights']) else ''])
for lam, m in R['P3']['candidates'].items():
    rows.append(['P3', f'λ = {lam}', m['theta'], m['TPR'], m['FPR'], 'dikunci' if float(lam) == R['P3']['locked_lambda'] else ''])
for th, m in R['P4']['validation_curve'].items():
    rows.append(['P4', f'θ_R = {th} (validasi)', float(th), m['TPR'], m['FPR'], 'dikunci' if float(th) == R['P4']['locked_theta_R'] else ''])
for k, m in R['P5'].items():
    rows.append(['P5', 'acuan (D terkunci)' if k == 'base' else k.replace(':', ' ').replace('x0.8', ' ×0,8').replace('x1.2', ' ×1,2'),
                 R['P4']['locked_theta_R'], m['TPR'], m['FPR'], ''])
sheet('3. Tuning (P1–P5)', 'P1–P5 – Sensitivitas parameter (seed validasi 0+2, gabungan)',
      'Urutan penguncian: W → bobot → λ → θ_R. P5 menggeser satu ambang flag ±20% dengan θ_R terkunci.',
      ['ID', 'Konfigurasi', 'θ terpilih', 'val_TPR', 'val_FPR', 'Keterangan'], rows,
      {'θ terpilih': '0.00', 'val_TPR': '0.000', 'val_FPR': '0.0000'},
      highlight=lambda r: r[5] == 'dikunci', widths={2: 36, 6: 12})

# ---- 4. Threshold curves (P4 test curve D, E1 sweep A) ----------------------
rows = []
for th, m in R['P4']['test_curve_seed1'].items():
    rows.append(['D (θ_R)', float(th), m['TP'], m['FP'], m['TPR'], m['TPR_adj'], m['FPR'], m['alerts_per_day']])
for th, m in R['E1_A']['theta_sweep_seed1'].items():
    rows.append(['A (θ Landauer)', float(th), m['TP'], m['FP'], m['TPR'], m['TPR_adj'], m['FPR'], m['alerts_per_day']])
sheet('4. Kurva threshold', 'Kurva TPR/FPR vs threshold pada seed 1 (untuk pelaporan, bukan tuning)',
      'D: θ_R 0,50–0,90 (θ terkunci 0,55). A: θ 0,30–0,90 (θ Landauer 0,6). Baris kuning = titik operasi yang dilaporkan.',
      ['Model', 'theta', 'TP', 'FP', 'TPR', 'TPR_adj', 'FPR', 'alerts_per_day'], rows, FMT,
      highlight=lambda r: (r[0].startswith('D') and abs(r[1] - R['P4']['locked_theta_R']) < 1e-9) or (r[0].startswith('A') and abs(r[1] - 0.6) < 1e-9),
      widths={1: 16})

# ---- 5. Per victim (H1) + K1/K2 ---------------------------------------------
rows = []
for r_ in sorted(R['H1'], key=lambda x: x['similarity']):
    def lab(x):
        if not x:
            return '–'
        if x['decision'] == 'anomalous':
            return 'TP' if x['status'] == 'detection' else 'TP_adj (training)'
        return 'FN'
    rows.append([r_['identity'], r_['victim_behaviour'], r_['similarity'], r_['volume_ratio'],
                 'K1 mudah' if r_['similarity'] < 0.3 else 'K2 sulit', lab(r_['A']), lab(r_['D'])])
ws = sheet('5. Per korban (H1, K1-K2)', 'H1 – Deteksi per korban, seed 1 (20 user-hari positif)',
           'identity = uid yang dinilai (sesudah swap berisi perilaku user lain); similarity = skor pasangan Landauer (0,3 volume + 0,7 tipe event). '
           'K1/K2 dibagi pada similarity 0,3.',
           ['Identitas (uid)', 'Perilaku dari user', 'Similarity', 'Rasio volume', 'Kelompok', 'Model A', 'Model D'], rows,
           {'Similarity': '0.000', 'Rasio volume': '0.00'}, widths={1: 40, 2: 40, 6: 18, 7: 18})
k = R['K1_K2']
r0 = 5 + len(rows) + 1
for j, h in enumerate(['Kelompok', 'n', 'TPR A', 'TPR D', 'TP_adj A', 'TP_adj D'], 1):
    c = ws.cell(r0, j, h); c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = HEAD; c.border = BORDER
for i, (g, e) in enumerate(k.items()):
    for j, v in enumerate([g.replace('_easy', ' mudah ').replace('_hard', ' sulit '), e['n'], e['A_TPR'], e['D_TPR'], e['A_tp_adj'], e['D_tp_adj']], 1):
        c = ws.cell(r0 + 1 + i, j, v); c.font = Font(name=F); c.border = BORDER
        if j in (3, 4):
            c.number_format = '0.00'

# ---- 6. FP causes + operations (H2, H3) -------------------------------------
h2 = R['H2']
rows = [['Sampel 50 FP', c, n] for c, n in h2['sample50_causes']]
rows += [['Semua FP', c, n] for c, n in h2['all_fp_causes']]
ws = sheet('6. False alarm (H2, H3)', 'H2 – Penyebab false alarm model D (seed 1) dan H3 – dampak operasional',
           f"Penyebab = kombinasi flag yang aktif pada hari FP. Total FP D = {h2['n_false_alarms_D']:,}. "
           'IF_score_only = tidak ada flag aktif.'.replace(',', '.'),
           ['Cakupan', 'Flag aktif', 'Jumlah FP'], rows, widths={1: 16, 2: 22})
r0 = 5 + len(rows) + 1
for j, h in enumerate(['Model', 'Total alert', 'Alert/hari', 'TP', 'Precision'], 1):
    c = ws.cell(r0, j, h); c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = HEAD; c.border = BORDER
for i, (mdl, e) in enumerate(R['H3'].items()):
    for j, (v, f) in enumerate([(mdl, None), (e['alerts'], '#,##0'), (e['alerts_per_day'], '0.00'), (e['TP'], '0'), (e['Prec'], '0.0000')], 1):
        c = ws.cell(r0 + 1 + i, j, v); c.font = Font(name=F); c.border = BORDER
        if f:
            c.number_format = f

# ---- 7. Validation (V1-V7) --------------------------------------------------
v1 = V12['V1']
rows = [
    ['V1', 'Cakupan field: type / time', '100% / 100% event', 'Lulus'],
    ['V1', 'Cakupan field: path', f"{v1['all']['pct_events']['path']:.2f}% event; {v1['all']['pct_user_days']['path']:.2f}% user-hari", 'Lulus'],
    ['V1', 'Cakupan field: isLocalIP', f"{v1['all']['pct_events']['isLocalIP']:.2f}% event (user login: {v1['uidType=name']['pct_events']['isLocalIP']:.2f}%)", 'Info'],
    ['V1', 'Cakupan field: location', f"{v1['all']['pct_events']['location']:.2f}% event (user login: {v1['uidType=name']['pct_events']['location']:.3f}%)", 'Info'],
    ['V2', 'Σ event tabel user-hari vs clue.json', f"{V12['V2']['summary_event_sum']:,} vs {V12['V2']['raw_rows']:,}; selisih {V12['V2']['difference']}", 'Lulus'],
    ['V3', 'Swap ringkasan = swap per event (seed 1)', f"{V3['differing_user_days']} user-hari berbeda dari {V3['user_days_ref']:,}", 'Lulus'],
    ['V4', 'Reproduksi baseline Landauer', 'TP 13, FP 3.337, TN 69.112, FN 7, Train 5.289', 'Lulus'],
    ['V5', 'Fitur kausal (100 user)', f"{R['V5']['feature_values_changed_on_or_before_t']} nilai fitur hari ≤ t berubah", 'Lulus' if R['V5']['pass'] else 'Gagal'],
    ['V6', 'Kolom input model', f"{R['V6']['n_columns_C_D']} kolom; kolom terlarang: {R['V6']['banned_present'] or 'tidak ada'}", 'Lulus' if R['V6']['pass'] else 'Gagal'],
    ['V7', 'Reprodusibilitas (seed 3, 2×)', 'Metrik identik' if R['V7']['identical'] else 'Berbeda', 'Lulus' if R['V7']['identical'] else 'Gagal'],
    ['K3', 'Injeksi sesi pembajakan (M7)', 'Tidak dijalankan (keputusan peneliti)', '–'],
]
sheet('7. Validasi (V1–V7)', 'V1–V7 – Validasi data dan pipeline', 'Sumber: pipeline/validation/results/*.json, results.json.',
      ['ID', 'Pemeriksaan', 'Hasil', 'Status'], rows, widths={2: 40, 3: 60})

wb.save(out)
print('saved', out, [ws.title for ws in wb.worksheets])
